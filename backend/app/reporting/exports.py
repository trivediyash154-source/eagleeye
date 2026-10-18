"""Structured exports: JSON (complete), CSV and XLSX (findings)."""
from __future__ import annotations

import csv
import io
import json

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy import select
from sqlalchemy.orm import Session

from ..core.util import iso, utcnow
from ..models import (
    Detection, EvidenceAsset, ExpectedComponent, ExpectedIdentifier, Finding, Inspection, InspectionVersion, OcrObservation, Report,
)
from ..services.common import audit, effective_verdict
from ..services.serialize import evidence_out, finding_out, inspection_out, review_out
from ..storage.local import get_storage

COLUMNS = ["seq", "check_type", "zone_code", "subject", "expected_value", "observed_value", "observed_raw", "verdict", "effective_verdict",
           "review_status", "rule", "ocr_confidence", "detector_confidence", "edit_distance", "review_reasons", "explanation", "finding_key", "finding_id"]


def _rows(findings: list[Finding]) -> list[list]:
    out = []
    for f in findings:
        c = f.confidences or {}
        out.append([f.seq, f.check_type, f.zone_code, f.subject, f.expected_value, f.observed_value, f.observed_raw, f.verdict, effective_verdict(f),
                    "reviewed" if f.reviews else "unreviewed", f.rule, c.get("ocr"), c.get("detector"), c.get("edit_distance"),
                    ";".join(f.review_reasons or []), f.explanation, f.finding_key, f.id])
    return out


def build_export(db: Session, version_id: str, fmt: str) -> Report:
    v = db.get(InspectionVersion, version_id)
    insp = db.get(Inspection, v.inspection_id)
    findings = db.scalars(select(Finding).where(Finding.version_id == version_id).order_by(Finding.seq)).all()
    if fmt == "json":
        doc = {
            "schema": "eagleeye.inspection-export/v1",
            "generated_at": iso(utcnow()),
            "disclaimer": "Supports inspection and human review; not a certificate of compliance.",
            "demonstration_fixture": insp.is_demo,
            "inspection": {k: val for k, val in inspection_out(db, insp).items() if k != "latest_version"},
            "version": {"id": v.id, "version_no": v.version_no, "label": v.label, "status": v.status, "config": v.config, "summary": v.summary,
                        "registration": v.registration, "preprocessing": v.preprocessing, "completed_at": iso(v.completed_at), "duration_ms": v.duration_ms},
            "inputs": [{"kind": a.kind, "filename": a.original_filename, "sha256": a.sha256, "size_bytes": a.size_bytes} for a in v.assets],
            "expected_components": [{"tag": c.tag, "zone": c.zone_code, "type": c.component_type, "class": c.canonical_class, "label": c.expected_label, "quantity": c.quantity,
                                     "source": f"{c.source_sheet}!{c.source_row}", "drawing_bbox": c.drawing_bbox, "drawing_match": c.drawing_match}
                                    for c in db.scalars(select(ExpectedComponent).where(ExpectedComponent.version_id == version_id))],
            "expected_identifiers": [{"kind": i.kind, "value": i.value, "terminal": i.terminal_ref, "wire_color": i.wire_color, "zone": i.zone_code,
                                      "source": f"{i.source_sheet}!{i.source_row}", "drawing_bbox": i.drawing_bbox, "drawing_match": i.drawing_match}
                                     for i in db.scalars(select(ExpectedIdentifier).where(ExpectedIdentifier.version_id == version_id))],
            "detections": [{"class": d.class_name, "confidence": d.confidence, "above_threshold": d.above_threshold, "bbox": [d.x0, d.y0, d.x1, d.y1], "model": d.model_name}
                           for d in db.scalars(select(Detection).where(Detection.version_id == version_id))],
            "ocr_observations": [{"raw": o.raw_text, "normalized": o.normalized_text, "confidence": o.confidence, "bbox": [o.x0, o.y0, o.x1, o.y1], "region": o.region, "engine": o.engine}
                                 for o in db.scalars(select(OcrObservation).where(OcrObservation.version_id == version_id))],
            "findings": [{**finding_out(f, v.version_no), "details": f.details, "reviews": [review_out(r) for r in f.reviews]} for f in findings],
            "evidence": [evidence_out(e) for e in db.scalars(select(EvidenceAsset).where(EvidenceAsset.version_id == version_id))],
        }
        data = json.dumps(doc, indent=2, default=str).encode()
        ext, name = "json", f"EAGLE-EYE_{insp.code}_v{v.version_no}.json"
    elif fmt == "csv":
        buf = io.StringIO()
        w = csv.writer(buf)
        w.writerow(COLUMNS)
        w.writerows(_rows(findings))
        data = buf.getvalue().encode("utf-8-sig")
        ext, name = "csv", f"EAGLE-EYE_{insp.code}_v{v.version_no}_findings.csv"
    elif fmt == "xlsx":
        wb = Workbook()
        ws = wb.active
        ws.title = "Findings"
        ws.append(COLUMNS)
        for r in _rows(findings):
            ws.append(r)
        fills = {"PASS": "D1FAE5", "FAIL": "FEE2E2", "UNCERTAIN": "FEF3C7", "WAIVED": "E5E7EB"}
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="0F172A")
        for row in ws.iter_rows(min_row=2):
            for idx in (7, 8):
                val = row[idx].value
                if val in fills:
                    row[idx].fill = PatternFill("solid", fgColor=fills[val])
            row[15].alignment = Alignment(wrap_text=True, vertical="top")
        widths = [5, 18, 7, 14, 16, 16, 16, 11, 13, 12, 26, 9, 9, 8, 26, 70, 26, 26]
        for i, wdt in enumerate(widths):
            ws.column_dimensions[chr(65 + i)].width = wdt
        ws.freeze_panes = "A2"
        s2 = wb.create_sheet("Summary")
        s2.append(["Inspection", insp.code])
        s2.append(["Cabinet", insp.cabinet_id])
        s2.append(["Version", v.version_no])
        s2.append(["Demonstration fixture", "YES — synthetic inputs" if insp.is_demo else "no"])
        for k in ("total", "pass", "fail", "uncertain"):
            s2.append([k.upper(), (v.summary or {}).get(k, 0)])
        s2.append(["Registration", (v.registration or {}).get("status")])
        s2.append(["Disclaimer", "Supports inspection and human review; not a certificate of compliance."])
        s2.column_dimensions["A"].width = 24
        s2.column_dimensions["B"].width = 60
        buf = io.BytesIO()
        wb.save(buf)
        data = buf.getvalue()
        ext, name = "xlsx", f"EAGLE-EYE_{insp.code}_v{v.version_no}_findings.xlsx"
    else:
        raise ValueError(f"unsupported export format {fmt}")
    st = get_storage()
    key = st.make_key("inspections", insp.id, v.id, "reports", ext=ext)
    sha, size = st.put_bytes(key, data)
    rep = Report(inspection_id=insp.id, version_id=v.id, format=fmt, storage_key=key, filename=name, size_bytes=size, sha256=sha)
    db.add(rep)
    db.flush()
    audit(db, "export_generated", "report", rep.id, insp.id, {"format": fmt, "version_no": v.version_no})
    db.flush()
    return rep
