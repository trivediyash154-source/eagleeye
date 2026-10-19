"""Generate the bundled SYNTHETIC demonstration fixtures for cabinet CAB-07.

Usage:  python scripts/generate_demo_fixtures.py  [--out ../data/demo]

Outputs (all computer-generated, labelled SYNTHETIC in their file names and content):
  CAB-07_cabinet_photo_v1_SYNTHETIC.jpg   initial state with deliberate discrepancies
  CAB-07_cabinet_photo_v2_SYNTHETIC.jpg   state after corrections (reinspection)
  CAB-07_GA_drawing_SYNTHETIC.pdf         vector GA drawing of the expected layout
  CAB-07_expected_list_SYNTHETIC.xlsx     component list + wire list (two sheets)
  CAB-07_wire_list_SYNTHETIC.csv          the wire list alone, as CSV
  fixtures.json                           manifest incl. the ground truth of the deliberate discrepancies
"""
from __future__ import annotations

import argparse
import csv
import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from app.demo.synth_cabinet import demo_layout, encode_jpeg, render_ga_pdf, render_photo  # noqa: E402

TYPE_TEXT = {"circuit_breaker": "Miniature circuit breaker (MCB)", "contactor": "Contactor", "power_supply": "Power supply 24 V DC", "earth_bar": "PE earth bar"}
COLOR_TEXT = {"brown": "Brown", "blue": "Blue", "black": "Black", "grey": "Grey"}


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", type=Path, default=Path(__file__).resolve().parents[2] / "data" / "demo")
    a = ap.parse_args()
    out: Path = a.out
    out.mkdir(parents=True, exist_ok=True)

    v1 = demo_layout("v1")
    v2 = demo_layout("v2")
    img1, _, _ = render_photo(v1, seed=7, ppm=5.0, out_size=(4000, 3000), jitter=0.035)
    (out / "CAB-07_cabinet_photo_v1_SYNTHETIC.jpg").write_bytes(encode_jpeg(img1, 93))
    img2, _, _ = render_photo(v2, seed=23, ppm=5.0, out_size=(4000, 3000), jitter=0.04)
    (out / "CAB-07_cabinet_photo_v2_SYNTHETIC.jpg").write_bytes(encode_jpeg(img2, 93))
    render_ga_pdf(v2, out / "CAB-07_GA_drawing_SYNTHETIC.pdf")

    wb = Workbook()
    cover = wb.active
    cover.title = "Components"
    cover.append(["EAGLE EYE demonstration fixture — SYNTHETIC expected component list for cabinet CAB-07"])
    cover.append(["Generated for software demonstration; not a real project document."])
    cover.append([])
    cover.append(["Zone", "Ref. Designation", "Component Type", "Device Label", "Qty", "Remarks"])
    for c in cover[4]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="0F172A")
    for it in v2.items:
        cover.append([it.zone, it.tag, TYPE_TEXT[it.kind], it.tag, 1, it.description])
    cover.append(["Z4", "X12", "Terminal block", "X12", len(v2.terminals), "Terminal strip X12, 10 terminals"])
    for col, w in zip("ABCDEF", (8, 18, 32, 14, 6, 30)):
        cover.column_dimensions[col].width = w
    ws = wb.create_sheet("Wire List")
    ws.append(["Terminal", "Ferrule No.", "Wire Colour", "From Device", "Zone", "Notes"])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="0F172A")
    rows = []
    for t in v2.terminals:
        row = [t.ref, t.ferrule_expected, COLOR_TEXT.get(t.wire_color_expected, t.wire_color_expected), t.source_tag, t.zone, ""]
        ws.append(row)
        rows.append(row)
    for col, w in zip("ABCDEF", (10, 14, 12, 12, 6, 20)):
        ws.column_dimensions[col].width = w
    wb.save(out / "CAB-07_expected_list_SYNTHETIC.xlsx")
    with open(out / "CAB-07_wire_list_SYNTHETIC.csv", "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["Terminal", "Ferrule No.", "Wire Colour", "From Device", "Zone", "Notes"])
        w.writerows(rows)

    manifest = {
        "fixture_set": "CAB-07",
        "synthetic": True,
        "notice": "All files are procedurally generated for demonstration and testing. They are not photographs or documents of a real cabinet.",
        "files": {
            "photo_v1": "CAB-07_cabinet_photo_v1_SYNTHETIC.jpg",
            "photo_v2": "CAB-07_cabinet_photo_v2_SYNTHETIC.jpg",
            "drawing": "CAB-07_GA_drawing_SYNTHETIC.pdf",
            "expected_list": "CAB-07_expected_list_SYNTHETIC.xlsx",
            "wire_list_csv": "CAB-07_wire_list_SYNTHETIC.csv",
        },
        "ground_truth_discrepancies_v1": [
            {"item": "CB44", "check": "component_presence", "truth": "breaker not fitted (label plate present)"},
            {"item": "PS12", "check": "device_label", "truth": "label plate printed 'PS21'"},
            {"item": "K2", "check": "device_label", "truth": "label plate partly covered by a cable tie"},
            {"item": "X12:4", "check": "ferrule_label", "truth": "ferrule 'X12-8B' fitted instead of 'X12-4B'"},
            {"item": "X12:6", "check": "ferrule_label", "truth": "ferrule printed 'X12-68' instead of 'X12-6B' (8/B confusable) — present in v1 and v2"},
            {"item": "X12:7", "check": "wire_color", "truth": "brown wire instead of blue (experimental check)"},
        ],
        "fixed_in_v2": ["CB44", "PS12", "X12:4", "X12:7"],
        "still_present_in_v2": ["K2 (label still covered)", "X12:6"],
    }
    (out / "fixtures.json").write_text(json.dumps(manifest, indent=2), encoding="utf-8")
    print("fixtures written to", out)


if __name__ == "__main__":
    main()
